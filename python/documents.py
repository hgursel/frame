"""Bounded offline extraction and document generation. JSON in, JSON out."""
import json
import os
import re
import resource
import sys
import tempfile
import zipfile
from pathlib import Path

resource.setrlimit(resource.RLIMIT_CPU, (35, 35))
resource.setrlimit(resource.RLIMIT_AS, (1024 * 1024 * 1024, 1024 * 1024 * 1024))
MAX_TEXT = 120_000


def status():
    from importlib.metadata import version
    expected = {"pypdf": "6.18.1", "reportlab": "5.0.1", "python-docx": "1.2.0"}
    return {"ready": all(version(name) == value for name, value in expected.items())}


def extract(data):
    source = Path(data["path"])
    if source.stat().st_size > 10 * 1024 * 1024:
        raise ValueError("File is too large")
    parts = []
    truncated = False
    if data["extension"] == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(source)
        if reader.is_encrypted:
            raise ValueError("Password-protected PDFs are not supported")
        for page in reader.pages[:100]:
            parts.append(page.extract_text() or "")
            if sum(map(len, parts)) > MAX_TEXT:
                truncated = True
                break
        truncated |= len(reader.pages) > 100
    elif data["extension"] == ".docx":
        with zipfile.ZipFile(source) as archive:
            if sum(i.file_size for i in archive.infolist()) > 64 * 1024 * 1024 or len(archive.infolist()) > 2000:
                raise ValueError("DOCX archive is too large when expanded")
            names = archive.namelist()
            if "word/document.xml" not in names or any("vbaProject" in n for n in names):
                raise ValueError("Only ordinary, non-macro DOCX documents are supported")
        from docx import Document
        doc = Document(source)
        parts.extend(p.text for p in doc.paragraphs)
        for table in doc.tables:
            parts.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
    else:
        raise ValueError("Unsupported extraction format")
    text = "\n\n".join(parts)
    if not text.strip():
        raise ValueError("No selectable text found. Scanned PDFs need OCR; OCR is not included.")
    return {"text": text[:MAX_TEXT], "truncated": truncated or len(text) > MAX_TEXT}


def blocks(markdown):
    # Deliberately small Markdown subset. Treat every line as text, never HTML/code.
    for line in markdown.splitlines():
        if line.startswith("### "):
            yield "h3", line[4:]
        elif line.startswith("## "):
            yield "h2", line[3:]
        elif line.startswith("# "):
            yield "h1", line[2:]
        elif re.match(r"^[-*] ", line):
            yield "bullet", line[2:]
        elif line.strip():
            yield "p", line


def generate(data):
    if data.get("format") != "docx":
        raise ValueError("PDF generation is available only through the Reports plugin. Use reports_create; enable Reports in system and project settings if needed.")
    title = str(data["title"])[:200]
    markdown = str(data["markdown"])
    if len(markdown) > MAX_TEXT:
        raise ValueError("Document content is too long")
    directory = Path(data["directory"]).resolve(strict=True)
    name = str(data["filename"])
    fmt = data["format"]
    if Path(name).name != name or not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]{0,100}", name):
        raise ValueError("Invalid output filename")
    output = directory / (name if name.endswith("." + fmt) else name + "." + fmt)
    descriptor, temporary = tempfile.mkstemp(prefix=".building-", dir=directory)
    os.close(descriptor)
    try:
        from docx import Document
        from docx.shared import Pt
        doc = Document()
        doc.styles["Normal"].font.name = "Calibri"
        doc.styles["Normal"].font.size = Pt(11)
        doc.add_heading(title, 0)
        for kind, text in blocks(markdown):
            if kind.startswith("h"):
                doc.add_heading(text, int(kind[1]))
            else:
                doc.add_paragraph(text, style="List Bullet" if kind == "bullet" else None)
        doc.save(temporary)
        # A complete file becomes visible atomically, without overwriting existing artifacts.
        os.link(temporary, output)
        os.unlink(temporary)
        return {"name": output.name, "bytes": output.stat().st_size}
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


try:
    request = json.loads(sys.stdin.read(1024 * 1024))
    operation = {"status": status, "extract": extract, "generate": generate}[request["command"]]
    result = operation() if request["command"] == "status" else operation(request)
    print(json.dumps(result))
except Exception as error:
    # Local diagnostics deliberately omit filenames, file contents, and arbitrary library errors.
    message = str(error) if isinstance(error, ValueError) else "Document operation failed; check dependencies and file format."
    print(json.dumps({"error": message}))
    sys.exit(1)
